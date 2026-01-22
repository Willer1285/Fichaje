"""
Generador de informes en PDF y Excel
Cumple con requisitos de Inspección de Trabajo
"""

from datetime import datetime
from typing import List, Tuple, Dict
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from app.database.models import Fichaje, Employee


class ReportGenerator:
    """Generador de informes"""

    @staticmethod
    def _clasificar_fichajes_por_tipo(fichajes: List[Fichaje]) -> Dict[str, float]:
        """Clasifica y suma las horas trabajadas por tipo de fichaje"""
        clasificacion = {
            "normal": 0.0,
            "feriado": 0.0,
            "fin_semana": 0.0,
            "horas_extra": 0.0,
            "salida_vacaciones": 0.0,
            "entrada_vacaciones": 0.0
        }

        for fichaje in fichajes:
            tipo = fichaje.tipo_fichaje or "normal"
            # Para horas extras, usar horas aprobadas en lugar de horas trabajadas
            if tipo == "horas_extra":
                horas = fichaje.horas_extras_aprobadas if fichaje.horas_extras_aprobadas > 0 else fichaje.horas_trabajadas
            else:
                horas = fichaje.horas_trabajadas
            if tipo in clasificacion:
                clasificacion[tipo] += horas
            else:
                clasificacion["normal"] += horas  # Por defecto

        return clasificacion

    @staticmethod
    def _get_nombre_tipo_fichaje(tipo: str) -> str:
        """Devuelve el nombre legible del tipo de fichaje"""
        nombres = {
            "normal": "Días normales",
            "feriado": "Días feriados",
            "fin_semana": "Fines de semana",
            "horas_extra": "Horas extras",
            "salida_vacaciones": "Salida de vacaciones",
            "entrada_vacaciones": "Entrada de vacaciones"
        }
        return nombres.get(tipo, tipo.capitalize())

    @staticmethod
    def generar_pdf_empleado(empleado: Employee, fichajes: List[Fichaje],
                             fecha_inicio: datetime, fecha_fin: datetime,
                             filename: str) -> str:
        """
        Genera un PDF con el registro de jornada de un empleado
        Formato válido para Inspección de Trabajo
        """
        doc = SimpleDocTemplate(filename, pagesize=A4)
        elements = []
        styles = getSampleStyleSheet()

        # Título
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#1976D2'),
            spaceAfter=30,
            alignment=1  # Center
        )

        title = Paragraph(
            f"REGISTRO DE JORNADA LABORAL<br/>Real Decreto-ley 8/2019",
            title_style
        )
        elements.append(title)
        elements.append(Spacer(1, 0.5 * cm))

        # Datos del empleado
        info_data = [
            ["Empleado:", f"{empleado.nombre} {empleado.apellidos}"],
            ["DNI/NIE:", empleado.dni],
            ["Nº Empleado:", empleado.numero_empleado or "N/A"],
            ["Tipo de Jornada:", empleado.tipo_jornada.capitalize()],
            ["Periodo:", f"{fecha_inicio.strftime('%d/%m/%Y')} - {fecha_fin.strftime('%d/%m/%Y')}"],
            ["Fecha de generación:", datetime.now().strftime('%d/%m/%Y %H:%M')]
        ]

        info_table = Table(info_data, colWidths=[4*cm, 12*cm])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#E3F2FD')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey)
        ]))

        elements.append(info_table)
        elements.append(Spacer(1, 1 * cm))

        # Tabla de fichajes (incluye tipo)
        fichajes_data = [["Fecha", "Tipo", "Entrada", "Salida Break", "Entrada Break", "Salida", "Horas"]]

        total_horas = 0
        for fichaje in sorted(fichajes, key=lambda x: x.fecha):
            tipo_abrev = {
                "normal": "Normal",
                "feriado": "Feriado",
                "fin_semana": "F.Semana",
                "horas_extra": "H.Extra",
                "salida_vacaciones": "S.Vac",
                "entrada_vacaciones": "E.Vac"
            }.get(fichaje.tipo_fichaje or "normal", "Normal")

            fichajes_data.append([
                fichaje.fecha.strftime('%d/%m/%Y') if fichaje.fecha else "",
                tipo_abrev,
                fichaje.hora_entrada.strftime('%H:%M') if fichaje.hora_entrada else "-",
                fichaje.hora_salida_break.strftime('%H:%M') if fichaje.hora_salida_break else "-",
                fichaje.hora_entrada_break.strftime('%H:%M') if fichaje.hora_entrada_break else "-",
                fichaje.hora_salida.strftime('%H:%M') if fichaje.hora_salida else "-",
                f"{fichaje.horas_trabajadas:.2f}h"
            ])
            total_horas += fichaje.horas_trabajadas

        # Fila de totales
        fichajes_data.append(["", "", "", "", "", "TOTAL:", f"{total_horas:.2f}h"])

        fichajes_table = Table(fichajes_data, colWidths=[2*cm, 2*cm, 2*cm, 2*cm, 2*cm, 2*cm, 1.5*cm])
        fichajes_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1976D2')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E3F2FD')),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ]))

        elements.append(fichajes_table)
        elements.append(Spacer(1, 1 * cm))

        # Resumen
        dias_trabajados = len(fichajes)
        promedio = total_horas / dias_trabajados if dias_trabajados > 0 else 0

        resumen_data = [
            ["Días trabajados:", str(dias_trabajados)],
            ["Total horas:", f"{total_horas:.2f}h"],
            ["Promedio diario:", f"{promedio:.2f}h"]
        ]

        resumen_table = Table(resumen_data, colWidths=[6*cm, 6*cm])
        resumen_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#E3F2FD')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey)
        ]))

        elements.append(resumen_table)
        elements.append(Spacer(1, 0.5 * cm))

        # Cálculo detallado adicional
        detalle_data = []

        # Calcular por semana si hay suficientes datos
        if len(fichajes) >= 7:
            semanas = (fecha_fin - fecha_inicio).days // 7
            if semanas > 0:
                horas_por_semana = total_horas / semanas
                detalle_data.append(["Horas promedio por semana:", f"{horas_por_semana:.2f}h"])

        # Clasificar horas por tipo
        clasificacion = ReportGenerator._clasificar_fichajes_por_tipo(fichajes)

        # Si tiene pago configurado, calcular con desglose
        if empleado.pago_por_hora > 0 or empleado.pago_hora_especial > 0:
            # Título de sección de cálculos
            if detalle_data:
                detalle_data.append(["", ""])  # Separador

            # Desglose de horas por tipo
            detalle_data.append(["DESGLOSE DE HORAS:", ""])
            for tipo, horas in clasificacion.items():
                if horas > 0:
                    nombre_tipo = ReportGenerator._get_nombre_tipo_fichaje(tipo)
                    detalle_data.append([f"  {nombre_tipo}:", f"{horas:.2f}h"])

            detalle_data.append(["", ""])  # Separador

            # Cálculo de pagos
            detalle_data.append(["CÁLCULO DE PAGOS:", ""])
            importe_total = 0.0

            # Horas normales (incluye feriados y fin de semana con tarifa normal si no hay tarifa extra)
            tarifa_normal = empleado.pago_por_hora
            horas_tarifa_normal = clasificacion["normal"]

            # Si no hay tarifa de hora extra, incluir feriados y fin de semana en tarifa normal
            if empleado.pago_hora_especial == 0:
                horas_tarifa_normal += clasificacion["feriado"] + clasificacion["fin_semana"]

            if horas_tarifa_normal > 0 and tarifa_normal > 0:
                importe_normal = horas_tarifa_normal * tarifa_normal
                detalle_data.append([f"  Horas normales ({horas_tarifa_normal:.2f}h x {tarifa_normal:.2f}€):",
                                   f"{importe_normal:.2f} €"])
                importe_total += importe_normal

            # Horas extras (incluye horas_extra, feriados y fin_semana si hay tarifa extra)
            if empleado.pago_hora_especial > 0:
                tarifa_extra = empleado.pago_hora_especial
                horas_extras = clasificacion["horas_extra"] + clasificacion["feriado"] + clasificacion["fin_semana"]

                if horas_extras > 0:
                    importe_extra = horas_extras * tarifa_extra
                    detalle_data.append([f"  Horas extras/especiales ({horas_extras:.2f}h x {tarifa_extra:.2f}€):",
                                       f"{importe_extra:.2f} €"])
                    importe_total += importe_extra

            detalle_data.append(["", ""])  # Separador
            detalle_data.append(["IMPORTE TOTAL:", f"{importe_total:.2f} €"])

        if detalle_data:
            detalle_table = Table(detalle_data, colWidths=[8*cm, 5*cm])
            detalle_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#FFF3E0')),
                ('ALIGN', (0, 0), (0, -1), 'LEFT'),
                ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                # Destacar las filas de título
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#FFE082')),
                # Destacar fila de total
                ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#FFE082')),
                ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
                ('FONTSIZE', (0, -1), (-1, -1), 12),
            ]))
            elements.append(detalle_table)
            elements.append(Spacer(1, 1 * cm))
        else:
            elements.append(Spacer(1, 1 * cm))

        # Sección de firma
        firma_style = ParagraphStyle(
            'Firma',
            parent=styles['Normal'],
            fontSize=10,
            spaceAfter=10
        )

        elements.append(Spacer(1, 1 * cm))
        elements.append(Paragraph("Certifico que los datos reflejados en este documento son correctos.", firma_style))
        elements.append(Spacer(1, 0.5 * cm))

        firma_data = [
            ["Fecha:", datetime.now().strftime('%d/%m/%Y'), "Firma del empleado:"],
            ["", "", ""],
            ["", "", "Firmo conforme"]
        ]

        firma_table = Table(firma_data, colWidths=[3*cm, 6*cm, 7*cm])
        firma_table.setStyle(TableStyle([
            ('ALIGN', (0, 0), (0, 0), 'LEFT'),
            ('ALIGN', (2, 0), (2, 0), 'LEFT'),
            ('FONTNAME', (0, 0), (0, 0), 'Helvetica-Bold'),
            ('FONTNAME', (2, 0), (2, 0), 'Helvetica-Bold'),
            ('FONTNAME', (2, 2), (2, 2), 'Helvetica-Oblique'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('LINEABOVE', (2, 2), (2, 2), 1, colors.black),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
        ]))

        elements.append(firma_table)
        elements.append(Spacer(1, 1.5 * cm))

        # Pie de página con info legal
        footer_style = ParagraphStyle(
            'Footer',
            parent=styles['Normal'],
            fontSize=8,
            textColor=colors.grey,
            alignment=1
        )

        footer = Paragraph(
            "Documento generado automáticamente conforme al Real Decreto-ley 8/2019<br/>"
            "Artículo 34.9 del Estatuto de los Trabajadores - Registro de Jornada Laboral<br/>"
            "Protección de datos según RGPD (UE) 2016/679",
            footer_style
        )
        elements.append(footer)

        # Generar PDF
        doc.build(elements)
        return filename

    @staticmethod
    def generar_excel_empleado(empleado: Employee, fichajes: List[Fichaje],
                               fecha_inicio: datetime, fecha_fin: datetime,
                               filename: str) -> str:
        """Genera un Excel con el registro de jornada de un empleado"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Registro de Jornada"

        # Estilos
        header_fill = PatternFill(start_color="1976D2", end_color="1976D2", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        info_fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
        bold_font = Font(bold=True)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Título
        ws.merge_cells('A1:F1')
        ws['A1'] = "REGISTRO DE JORNADA LABORAL - Real Decreto-ley 8/2019"
        ws['A1'].font = Font(bold=True, size=14, color="1976D2")
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

        # Información del empleado
        row = 3
        info = [
            ("Empleado:", f"{empleado.nombre} {empleado.apellidos}"),
            ("DNI/NIE:", empleado.dni),
            ("Nº Empleado:", empleado.numero_empleado or "N/A"),
            ("Tipo de Jornada:", empleado.tipo_jornada.capitalize()),
            ("Periodo:", f"{fecha_inicio.strftime('%d/%m/%Y')} - {fecha_fin.strftime('%d/%m/%Y')}"),
            ("Fecha de generación:", datetime.now().strftime('%d/%m/%Y %H:%M'))
        ]

        for label, value in info:
            ws[f'A{row}'] = label
            ws[f'A{row}'].font = bold_font
            ws[f'A{row}'].fill = info_fill
            ws[f'B{row}'] = value
            ws[f'A{row}'].border = border
            ws[f'B{row}'].border = border
            row += 1

        row += 1

        # Encabezados de tabla (incluye Tipo)
        headers = ["Fecha", "Tipo", "Entrada", "Salida Break", "Entrada Break", "Salida", "Horas"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

        row += 1

        # Datos de fichajes
        total_horas = 0
        for fichaje in sorted(fichajes, key=lambda x: x.fecha):
            tipo_abrev = {
                "normal": "Normal",
                "feriado": "Feriado",
                "fin_semana": "F.Semana",
                "horas_extra": "H.Extra",
                "salida_vacaciones": "S.Vac",
                "entrada_vacaciones": "E.Vac"
            }.get(fichaje.tipo_fichaje or "normal", "Normal")

            ws.cell(row=row, column=1, value=fichaje.fecha.strftime('%d/%m/%Y') if fichaje.fecha else "")
            ws.cell(row=row, column=2, value=tipo_abrev)
            ws.cell(row=row, column=3, value=fichaje.hora_entrada.strftime('%H:%M') if fichaje.hora_entrada else "-")
            ws.cell(row=row, column=4, value=fichaje.hora_salida_break.strftime('%H:%M') if fichaje.hora_salida_break else "-")
            ws.cell(row=row, column=5, value=fichaje.hora_entrada_break.strftime('%H:%M') if fichaje.hora_entrada_break else "-")
            ws.cell(row=row, column=6, value=fichaje.hora_salida.strftime('%H:%M') if fichaje.hora_salida else "-")
            ws.cell(row=row, column=7, value=f"{fichaje.horas_trabajadas:.2f}h")

            for col in range(1, 8):
                ws.cell(row=row, column=col).alignment = Alignment(horizontal='center')
                ws.cell(row=row, column=col).border = border

            total_horas += fichaje.horas_trabajadas
            row += 1

        # Total
        ws.cell(row=row, column=6, value="TOTAL:")
        ws.cell(row=row, column=6).font = bold_font
        ws.cell(row=row, column=7, value=f"{total_horas:.2f}h")
        ws.cell(row=row, column=7).font = bold_font
        for col in range(1, 8):
            ws.cell(row=row, column=col).fill = info_fill
            ws.cell(row=row, column=col).border = border

        row += 2

        # Resumen
        dias_trabajados = len(fichajes)
        promedio = total_horas / dias_trabajados if dias_trabajados > 0 else 0

        ws[f'A{row}'] = "Días trabajados:"
        ws[f'A{row}'].font = bold_font
        ws[f'B{row}'] = dias_trabajados

        row += 1
        ws[f'A{row}'] = "Total horas:"
        ws[f'A{row}'].font = bold_font
        ws[f'B{row}'] = f"{total_horas:.2f}h"

        row += 1
        ws[f'A{row}'] = "Promedio diario:"
        ws[f'A{row}'].font = bold_font
        ws[f'B{row}'] = f"{promedio:.2f}h"

        # Cálculo detallado adicional
        row += 2

        # Calcular por semana si hay suficientes datos
        if len(fichajes) >= 7:
            semanas = (fecha_fin - fecha_inicio).days // 7
            if semanas > 0:
                horas_por_semana = total_horas / semanas
                ws[f'A{row}'] = "Horas promedio por semana:"
                ws[f'A{row}'].font = bold_font
                ws[f'A{row}'].fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
                ws[f'B{row}'] = f"{horas_por_semana:.2f}h"
                ws[f'B{row}'].fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
                row += 1

        # Clasificar horas por tipo
        clasificacion = ReportGenerator._clasificar_fichajes_por_tipo(fichajes)

        # Si tiene pago configurado, calcular con desglose
        if empleado.pago_por_hora > 0 or empleado.pago_hora_especial > 0:
            row += 1
            title_fill = PatternFill(start_color="FFE082", end_color="FFE082", fill_type="solid")
            detalle_fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")

            # Desglose de horas por tipo
            ws[f'A{row}'] = "DESGLOSE DE HORAS:"
            ws[f'A{row}'].font = Font(bold=True, size=12)
            ws[f'A{row}'].fill = title_fill
            ws[f'B{row}'].fill = title_fill
            row += 1

            for tipo, horas in clasificacion.items():
                if horas > 0:
                    nombre_tipo = ReportGenerator._get_nombre_tipo_fichaje(tipo)
                    ws[f'A{row}'] = f"  {nombre_tipo}:"
                    ws[f'A{row}'].font = bold_font
                    ws[f'A{row}'].fill = detalle_fill
                    ws[f'B{row}'] = f"{horas:.2f}h"
                    ws[f'B{row}'].fill = detalle_fill
                    row += 1

            row += 1

            # Cálculo de pagos
            ws[f'A{row}'] = "CÁLCULO DE PAGOS:"
            ws[f'A{row}'].font = Font(bold=True, size=12)
            ws[f'A{row}'].fill = title_fill
            ws[f'B{row}'].fill = title_fill
            row += 1

            importe_total = 0.0

            # Horas normales
            tarifa_normal = empleado.pago_por_hora
            horas_tarifa_normal = clasificacion["normal"]

            # Si no hay tarifa de hora extra, incluir feriados y fin de semana
            if empleado.pago_hora_especial == 0:
                horas_tarifa_normal += clasificacion["feriado"] + clasificacion["fin_semana"]

            if horas_tarifa_normal > 0 and tarifa_normal > 0:
                importe_normal = horas_tarifa_normal * tarifa_normal
                ws[f'A{row}'] = f"  Horas normales ({horas_tarifa_normal:.2f}h x {tarifa_normal:.2f}€):"
                ws[f'A{row}'].font = bold_font
                ws[f'A{row}'].fill = detalle_fill
                ws[f'B{row}'] = f"{importe_normal:.2f} €"
                ws[f'B{row}'].fill = detalle_fill
                importe_total += importe_normal
                row += 1

            # Horas extras
            if empleado.pago_hora_especial > 0:
                tarifa_extra = empleado.pago_hora_especial
                horas_extras = clasificacion["horas_extra"] + clasificacion["feriado"] + clasificacion["fin_semana"]

                if horas_extras > 0:
                    importe_extra = horas_extras * tarifa_extra
                    ws[f'A{row}'] = f"  Horas extras/especiales ({horas_extras:.2f}h x {tarifa_extra:.2f}€):"
                    ws[f'A{row}'].font = bold_font
                    ws[f'A{row}'].fill = detalle_fill
                    ws[f'B{row}'] = f"{importe_extra:.2f} €"
                    ws[f'B{row}'].fill = detalle_fill
                    importe_total += importe_extra
                    row += 1

            row += 1
            ws[f'A{row}'] = "IMPORTE TOTAL:"
            ws[f'A{row}'].font = Font(bold=True, size=12)
            ws[f'A{row}'].fill = title_fill
            ws[f'B{row}'] = f"{importe_total:.2f} €"
            ws[f'B{row}'].font = Font(bold=True, size=12)
            ws[f'B{row}'].fill = title_fill
            row += 1

        # Sección de firma
        row += 3
        ws.merge_cells(f'A{row}:F{row}')
        ws[f'A{row}'] = "Certifico que los datos reflejados en este documento son correctos."
        ws[f'A{row}'].alignment = Alignment(horizontal='center')
        ws[f'A{row}'].font = Font(italic=True)

        row += 2
        ws[f'A{row}'] = f"Fecha: {datetime.now().strftime('%d/%m/%Y')}"
        ws[f'A{row}'].font = bold_font

        ws[f'D{row}'] = "Firma del empleado:"
        ws[f'D{row}'].font = bold_font

        row += 3
        ws[f'D{row}'] = "_____________________________"
        ws[f'D{row}'].alignment = Alignment(horizontal='center')

        row += 1
        ws[f'D{row}'] = "Firmo conforme"
        ws[f'D{row}'].alignment = Alignment(horizontal='center')
        ws[f'D{row}'].font = Font(italic=True)

        # Ajustar anchos de columna
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 12  # Tipo
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 14
        ws.column_dimensions['E'].width = 14
        ws.column_dimensions['F'].width = 12
        ws.column_dimensions['G'].width = 10

        # Guardar
        wb.save(filename)
        return filename

    @staticmethod
    def generar_pdf_todos(fichajes_empleados: List[Tuple[Fichaje, Employee]],
                          fecha_inicio: datetime, fecha_fin: datetime,
                          filename: str) -> str:
        """Genera un PDF con todos los fichajes del periodo (Lista plana)"""
        doc = SimpleDocTemplate(filename, pagesize=A4)
        elements = []
        styles = getSampleStyleSheet()

        # Título
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#1976D2'),
            spaceAfter=20,
            alignment=1  # Center
        )

        title = Paragraph(
            f"REGISTRO DE JORNADA - TODOS LOS EMPLEADOS<br/>{fecha_inicio.strftime('%d/%m/%Y')} al {fecha_fin.strftime('%d/%m/%Y')}",
            title_style
        )
        elements.append(title)
        elements.append(Spacer(1, 0.5 * cm))

        # Tabla de datos
        data = [["Fecha", "Empleado", "DNI", "Entrada", "Salida", "Horas"]]

        for fichaje, empleado in fichajes_empleados:
            data.append([
                fichaje.fecha.strftime('%d/%m/%Y') if fichaje.fecha else "",
                f"{empleado.apellidos}, {empleado.nombre}"[:30], # Truncar si es muy largo
                empleado.dni,
                fichaje.hora_entrada.strftime('%H:%M') if fichaje.hora_entrada else "-",
                fichaje.hora_salida.strftime('%H:%M') if fichaje.hora_salida else "-",
                f"{fichaje.horas_trabajadas:.2f}h"
            ])

        # Si no hay datos, añadir fila vacía o mensaje
        if len(data) == 1:
            data.append(["No hay registros", "", "", "", "", ""])

        # Estilo de tabla
        table = Table(data, colWidths=[2.5*cm, 7*cm, 2.5*cm, 2*cm, 2*cm, 2*cm])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1976D2')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('ALIGN', (1, 0), (1, -1), 'LEFT'), # Alinear nombres a la izquierda
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#E3F2FD')]),
        ]))

        elements.append(table)

        # Pie de página
        footer_style = ParagraphStyle(
            'Footer',
            parent=styles['Normal'],
            fontSize=8,
            textColor=colors.grey,
            alignment=1,
            spaceBefore=20
        )
        
        elements.append(Spacer(1, 1 * cm))
        footer = Paragraph(
            f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            footer_style
        )
        elements.append(footer)

        doc.build(elements)
        return filename

    @staticmethod
    def generar_excel_todos(fichajes_empleados: List[Tuple[Fichaje, Employee]],
                           fecha_inicio: datetime, fecha_fin: datetime,
                           filename: str) -> str:
        """Genera un Excel con todos los fichajes del periodo"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Todos los Fichajes"

        # Estilos
        header_fill = PatternFill(start_color="1976D2", end_color="1976D2", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Título
        ws.merge_cells('A1:H1')
        ws['A1'] = f"REGISTRO DE JORNADA LABORAL - {fecha_inicio.strftime('%d/%m/%Y')} al {fecha_fin.strftime('%d/%m/%Y')}"
        ws['A1'].font = Font(bold=True, size=14, color="1976D2")
        ws['A1'].alignment = Alignment(horizontal='center')

        # Encabezados
        headers = ["Fecha", "Empleado", "DNI", "Entrada", "Salida Break", "Entrada Break", "Salida", "Horas"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
            cell.border = border

        # Datos
        row = 4
        for fichaje, empleado in fichajes_empleados:
            ws.cell(row=row, column=1, value=fichaje.fecha.strftime('%d/%m/%Y') if fichaje.fecha else "")
            ws.cell(row=row, column=2, value=f"{empleado.apellidos}, {empleado.nombre}")
            ws.cell(row=row, column=3, value=empleado.dni)
            ws.cell(row=row, column=4, value=fichaje.hora_entrada.strftime('%H:%M') if fichaje.hora_entrada else "-")
            ws.cell(row=row, column=5, value=fichaje.hora_salida_break.strftime('%H:%M') if fichaje.hora_salida_break else "-")
            ws.cell(row=row, column=6, value=fichaje.hora_entrada_break.strftime('%H:%M') if fichaje.hora_entrada_break else "-")
            ws.cell(row=row, column=7, value=fichaje.hora_salida.strftime('%H:%M') if fichaje.hora_salida else "-")
            ws.cell(row=row, column=8, value=f"{fichaje.horas_trabajadas:.2f}h")

            for col in range(1, 9):
                ws.cell(row=row, column=col).alignment = Alignment(horizontal='center')
                ws.cell(row=row, column=col).border = border

            row += 1

        # Ajustar anchos
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 14
        ws.column_dimensions['F'].width = 14
        ws.column_dimensions['G'].width = 12
        ws.column_dimensions['H'].width = 10

        wb.save(filename)
        return filename
